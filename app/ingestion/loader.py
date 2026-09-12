from pathlib import Path

import pymupdf as fitz  # PyMuPDF

from langchain_core.documents import Document
from app.config import OCR_ENABLED
from app.logger import logger


class DocumentLoader:
    """Loads supported documents: PDF (with OCR fallback), DOCX, TXT, MD."""

    SUPPORTED_SUFFIXES: frozenset = frozenset({".pdf", ".docx", ".txt", ".md"})

    # If a PDF's pages yield less text than this, treat it as scanned
    MIN_TEXT_CHARS = 50

    @classmethod
    def load(cls, file_path: str) -> list[Document]:
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(file_path)

        suffix = path.suffix.lower()

        if suffix not in cls.SUPPORTED_SUFFIXES:
            raise ValueError(
                f"Unsupported file type '{suffix}'. "
                f"Supported: {sorted(cls.SUPPORTED_SUFFIXES)}"
            )

        if suffix == ".pdf":
            return cls._load_pdf(path)
        if suffix == ".docx":
            return cls._load_docx(path)
        return cls._load_plaintext(path)

    @classmethod
    def _load_pdf(cls, path: Path) -> list[Document]:
        docs = cls._pdf_with_pymupdf(path)

        total_chars = sum(len(d.page_content.strip()) for d in docs)
        if total_chars >= cls.MIN_TEXT_CHARS:
            return docs

        # Image-only PDF: no embedded text layer
        extractable = total_chars > 0
        logger.info(
            f"{path.name}: {'little' if extractable else 'no'} extractable text "
            f"({total_chars} chars)."
        )

        if not OCR_ENABLED:
            if extractable:
                return docs
            raise ValueError(
                f"'{path.name}' looks like a scanned (image-only) PDF with no "
                "text layer, and OCR is disabled. Enable OCR_ENABLED=true."
            )

        ocr_docs = cls._ocr_pdf(path)
        if not ocr_docs:
            if extractable:
                return docs
            raise ValueError(
                f"OCR produced no text for '{path.name}'. "
                "Is tesseract installed?"
            )
        logger.success(f"{path.name}: OCR extracted {len(ocr_docs)} pages")
        return ocr_docs

    @staticmethod
    def _pdf_with_pymupdf(path: Path) -> list[Document]:
        from langchain_community.document_loaders import PyMuPDFLoader

        logger.info(f"Loading {path}")
        return PyMuPDFLoader(str(path)).load()

    @staticmethod
    def _ocr_pdf(path: Path) -> list[Document]:
        """OCR a scanned PDF page-by-page via PyMuPDF rendering + tesseract."""

        docs: list[Document] = []
        pdf = fitz.open(path)

        try:
            for i, page in enumerate(pdf):
                pix = page.get_pixmap(dpi=200)
                text = _ocr_pixmap(pix)

                if text.strip():
                    docs.append(
                        Document(
                            page_content=text,
                            metadata={"source": str(path), "page": i},
                        )
                    )
        finally:
            pdf.close()

        return docs

    @staticmethod
    def _load_docx(path: Path) -> list[Document]:
        import docx as docx_lib

        doc = docx_lib.Document(str(path))
        text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

        if not text.strip():
            raise ValueError(f"'{path.name}' contains no text content.")

        return [
            Document(
                page_content=text,
                metadata={"source": str(path), "page": 0},
            )
        ]

    @staticmethod
    def _load_plaintext(path: Path) -> list[Document]:
        text = path.read_text(encoding="utf-8", errors="replace")

        if not text.strip():
            raise ValueError(f"'{path.name}' is empty.")

        return [
            Document(
                page_content=text,
                metadata={"source": str(path), "page": 0},
            )
        ]


def _ocr_pixmap(pix) -> str:
    """Run tesseract OCR on a PyMuPDF pixmap via PIL."""

    import pytesseract
    from PIL import Image

    mode = "RGBA" if pix.alpha else "RGB"
    img = Image.frombytes(mode, (pix.width, pix.height), pix.samples)
    if mode == "RGBA":
        img = img.convert("RGB")

    return pytesseract.image_to_string(img)
